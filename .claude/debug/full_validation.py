#!/usr/bin/env python3
"""
Validación completa Frontend vs Backend con Playwright y Excel.
- Lanza servidor HTTP
- Abre el visualizador
- Carga el JSON
- Navega por cada fase
- Compara valores UI vs JSON
- Captura screenshots
- Genera Excel con resultados
"""

import json
import subprocess
import time
import os
import sys
from pathlib import Path
from datetime import datetime

from playwright.sync_api import sync_playwright, Page
from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


class FullValidator:
    MATERIALS = ['cereal', 'mineral', 'clay', 'wood', 'wool']

    def __init__(self, visualizer_path: str, trace_path: str, output_dir: str):
        self.visualizer_path = Path(visualizer_path)
        self.trace_path = Path(trace_path)
        self.output_dir = Path(output_dir)
        self.output_dir.mkdir(parents=True, exist_ok=True)

        self.server_process = None
        self.page = None
        self.trace_data = None

        self.results = []  # Lista de comparaciones
        self.screenshots = []  # Lista de paths de screenshots
        self.errors = []

    def load_trace(self):
        """Carga el JSON de la traza."""
        with open(self.trace_path) as f:
            self.trace_data = json.load(f)
        print(f"Traza cargada: {self.trace_path}")

    def start_server(self, port=8765):
        """Lanza servidor HTTP."""
        self.server_process = subprocess.Popen(
            ['python3', '-m', 'http.server', str(port)],
            cwd=self.visualizer_path,
            stdout=subprocess.DEVNULL,
            stderr=subprocess.DEVNULL
        )
        time.sleep(1)
        print(f"Servidor HTTP lanzado en puerto {port}")
        return f"http://localhost:{port}"

    def stop_server(self):
        """Detiene el servidor HTTP."""
        if self.server_process:
            self.server_process.terminate()
            self.server_process.wait()
            print("Servidor HTTP detenido")

    def parse_json_hand(self, hand):
        """Parsea hand del JSON (maneja strings e ints)."""
        if not hand:
            return {m: 0 for m in self.MATERIALS}
        result = {}
        for m in self.MATERIALS:
            val = hand.get(m, 0)
            result[m] = int(val) if isinstance(val, str) else val
        return result

    def get_ui_resources(self, player: int) -> dict:
        """Obtiene recursos mostrados en UI."""
        resources = {}
        for mat in self.MATERIALS:
            selector = f'#hand_P{player} .{mat} .{mat}_quantity'
            try:
                element = self.page.query_selector(selector)
                if element:
                    text = element.inner_text().strip()
                    resources[mat] = int(text) if text else 0
                else:
                    resources[mat] = 0
            except Exception as e:
                resources[mat] = -1  # Error
        return resources

    def take_screenshot(self, name: str) -> str:
        """Captura screenshot y retorna el path."""
        filename = f"{name}.png"
        filepath = self.output_dir / filename
        self.page.screenshot(path=str(filepath))
        self.screenshots.append(str(filepath))
        return str(filepath)

    def compare_and_record(self, round_num: int, turn: int, phase: str,
                          player: int, expected: dict, actual: dict, screenshot_path: str):
        """Compara y registra resultado."""
        match = True
        differences = []

        for mat in self.MATERIALS:
            exp = expected.get(mat, 0)
            act = actual.get(mat, 0)
            if exp != act:
                match = False
                differences.append(f"{mat}: esperado {exp}, actual {act}")

        result = {
            'round': round_num,
            'turn': turn,
            'phase': phase,
            'player': player,
            'expected': expected.copy(),
            'actual': actual.copy(),
            'match': match,
            'differences': differences,
            'screenshot': screenshot_path
        }
        self.results.append(result)

        if not match:
            self.errors.append(result)
            print(f"  ❌ P{player} MISMATCH: {', '.join(differences)}")
        else:
            print(f"  ✓ P{player} OK")

        return match

    def validate_phase(self, round_num: int, turn: int, phase_key: str, phase_data: dict):
        """Valida una fase específica."""
        context = f"R{round_num}_T{turn}_{phase_key}"
        print(f"\nValidando {context}...")

        # Tomar screenshot
        screenshot_path = self.take_screenshot(context)

        # Comparar recursos de cada jugador
        for player in range(4):
            hand_key = f'hand_P{player}'
            if hand_key in phase_data:
                expected = self.parse_json_hand(phase_data[hand_key])
                actual = self.get_ui_resources(player)
                self.compare_and_record(round_num, turn, phase_key, player,
                                       expected, actual, screenshot_path)

    def navigate_to_round(self, round_num: int):
        """Navega a una ronda específica."""
        try:
            selector = self.page.query_selector('#contador_rondas')
            if selector:
                selector.select_option(f'round_{round_num}')
                time.sleep(0.3)
                return True
        except Exception as e:
            print(f"Error navegando a ronda {round_num}: {e}")
        return False

    def navigate_to_turn(self, turn: int):
        """Navega a un turno específico."""
        try:
            selector = self.page.query_selector('#contador_turnos')
            if selector:
                selector.select_option(f'P{turn}')
                time.sleep(0.3)
                return True
        except Exception as e:
            print(f"Error navegando a turno {turn}: {e}")
        return False

    def navigate_to_phase(self, phase: str):
        """Navega a una fase específica."""
        try:
            selector = self.page.query_selector('#contador_fases')
            if selector:
                # Intentar diferentes formatos
                for option in [phase, f'{phase}_0', phase.replace('_', '')]:
                    try:
                        selector.select_option(option)
                        time.sleep(0.3)
                        return True
                    except:
                        continue
        except Exception as e:
            print(f"Error navegando a fase {phase}: {e}")
        return False

    def run_validation(self, base_url: str, max_rounds: int = 3):
        """Ejecuta la validación completa."""
        print(f"\n{'='*60}")
        print("INICIANDO VALIDACIÓN COMPLETA")
        print(f"{'='*60}")

        with sync_playwright() as p:
            browser = p.chromium.launch(headless=True)
            self.page = browser.new_page()

            # Cargar visualizador
            print(f"\nCargando visualizador: {base_url}")
            self.page.goto(base_url)
            time.sleep(2)

            # Tomar screenshot inicial
            self.take_screenshot("00_initial")

            # Cargar JSON
            print(f"Cargando JSON: {self.trace_path}")
            file_input = self.page.query_selector('input[type="file"]')
            if file_input:
                file_input.set_input_files(str(self.trace_path))
                time.sleep(2)
                self.take_screenshot("01_json_loaded")
            else:
                print("ERROR: No se encontró input de archivo")
                browser.close()
                return

            # Validar cada ronda
            game = self.trace_data.get('game', {})
            round_keys = sorted([k for k in game.keys() if k.startswith('round_')])[:max_rounds]

            for round_key in round_keys:
                round_num = int(round_key.replace('round_', ''))
                round_obj = game[round_key]

                print(f"\n{'='*40}")
                print(f"RONDA {round_num}")
                print(f"{'='*40}")

                self.navigate_to_round(round_num)

                # Validar cada turno
                for turn in range(4):
                    turn_key = f'turn_P{turn}'
                    if turn_key not in round_obj:
                        continue

                    turn_obj = round_obj[turn_key]
                    self.navigate_to_turn(turn)

                    # Validar start_turn
                    if 'start_turn' in turn_obj:
                        self.navigate_to_phase('start_turn')
                        self.validate_phase(round_num, turn, 'start_turn', turn_obj['start_turn'])

                    # Validar end_turn
                    if 'end_turn' in turn_obj:
                        self.navigate_to_phase('end_turn')
                        self.validate_phase(round_num, turn, 'end_turn', turn_obj['end_turn'])

            browser.close()

    def generate_excel(self, output_path: str):
        """Genera Excel con todos los resultados."""
        wb = Workbook()

        # Hoja de Resumen
        ws_summary = wb.active
        ws_summary.title = "Resumen"

        # Estilos
        header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        header_font = Font(bold=True, color="FFFFFF")
        ok_fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
        error_fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")

        # Encabezado resumen
        ws_summary['A1'] = "VALIDACIÓN FRONTEND vs BACKEND"
        ws_summary['A1'].font = Font(bold=True, size=14)
        ws_summary['A3'] = f"Archivo: {self.trace_path}"
        ws_summary['A4'] = f"Fecha: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}"
        ws_summary['A6'] = f"Total comparaciones: {len(self.results)}"
        ws_summary['A7'] = f"Coincidencias: {len([r for r in self.results if r['match']])}"
        ws_summary['A8'] = f"Discrepancias: {len(self.errors)}"

        if self.errors:
            ws_summary['A8'].fill = error_fill
        else:
            ws_summary['A8'].fill = ok_fill

        # Hoja de Detalle
        ws_detail = wb.create_sheet("Detalle")
        headers = ['Ronda', 'Turno', 'Fase', 'Jugador',
                   'Cereal Exp', 'Cereal Act',
                   'Mineral Exp', 'Mineral Act',
                   'Clay Exp', 'Clay Act',
                   'Wood Exp', 'Wood Act',
                   'Wool Exp', 'Wool Act',
                   'Match', 'Diferencias']

        for col, header in enumerate(headers, 1):
            cell = ws_detail.cell(row=1, column=col, value=header)
            cell.fill = header_fill
            cell.font = header_font

        for row_idx, result in enumerate(self.results, 2):
            ws_detail.cell(row=row_idx, column=1, value=result['round'])
            ws_detail.cell(row=row_idx, column=2, value=result['turn'])
            ws_detail.cell(row=row_idx, column=3, value=result['phase'])
            ws_detail.cell(row=row_idx, column=4, value=result['player'])

            col = 5
            for mat in self.MATERIALS:
                ws_detail.cell(row=row_idx, column=col, value=result['expected'].get(mat, 0))
                ws_detail.cell(row=row_idx, column=col+1, value=result['actual'].get(mat, 0))
                col += 2

            match_cell = ws_detail.cell(row=row_idx, column=15, value="✓" if result['match'] else "✗")
            match_cell.fill = ok_fill if result['match'] else error_fill

            ws_detail.cell(row=row_idx, column=16, value='; '.join(result['differences']))

        # Hoja de Errores
        ws_errors = wb.create_sheet("Errores")
        error_headers = ['Ronda', 'Turno', 'Fase', 'Jugador', 'Diferencias', 'Screenshot']

        for col, header in enumerate(error_headers, 1):
            cell = ws_errors.cell(row=1, column=col, value=header)
            cell.fill = PatternFill(start_color="FF0000", end_color="FF0000", fill_type="solid")
            cell.font = header_font

        if not self.errors:
            ws_errors.cell(row=2, column=1, value="Sin errores detectados")
            ws_errors['A2'].fill = ok_fill
        else:
            for row_idx, error in enumerate(self.errors, 2):
                ws_errors.cell(row=row_idx, column=1, value=error['round'])
                ws_errors.cell(row=row_idx, column=2, value=error['turn'])
                ws_errors.cell(row=row_idx, column=3, value=error['phase'])
                ws_errors.cell(row=row_idx, column=4, value=error['player'])
                ws_errors.cell(row=row_idx, column=5, value='; '.join(error['differences']))
                ws_errors.cell(row=row_idx, column=6, value=error['screenshot'])

        # Ajustar anchos de columna
        for ws in [ws_summary, ws_detail, ws_errors]:
            for col in range(1, 20):
                ws.column_dimensions[get_column_letter(col)].width = 15

        wb.save(output_path)
        print(f"\nExcel generado: {output_path}")

    def run(self, port=8765, max_rounds=3):
        """Ejecuta todo el proceso de validación."""
        try:
            self.load_trace()
            base_url = self.start_server(port)
            self.run_validation(base_url, max_rounds)

            excel_path = self.output_dir / "validacion_resultado.xlsx"
            self.generate_excel(str(excel_path))

            # Resumen final
            print(f"\n{'='*60}")
            print("RESUMEN FINAL")
            print(f"{'='*60}")
            print(f"Total comparaciones: {len(self.results)}")
            print(f"Coincidencias: {len([r for r in self.results if r['match']])}")
            print(f"Discrepancias: {len(self.errors)}")
            print(f"Screenshots: {len(self.screenshots)}")
            print(f"Excel: {excel_path}")
            print(f"{'='*60}")

            return len(self.errors) == 0

        finally:
            self.stop_server()


def main():
    if len(sys.argv) < 3:
        print("Uso: python full_validation.py <visualizer_path> <trace.json> [output_dir]")
        print("Ejemplo: python full_validation.py ./Visualizer ./trace.json ./output")
        sys.exit(1)

    visualizer_path = sys.argv[1]
    trace_path = sys.argv[2]
    output_dir = sys.argv[3] if len(sys.argv) > 3 else "./.claude/logs/debug-sessions"

    validator = FullValidator(visualizer_path, trace_path, output_dir)
    success = validator.run(max_rounds=3)

    sys.exit(0 if success else 1)


if __name__ == '__main__':
    main()
